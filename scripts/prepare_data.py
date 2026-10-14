"""Build the assignment's small, traceable data files from the local source archive.

Usage: python scripts/prepare_data.py --source-dir '../Local Tourism Attractions Data'
Original downloads are never modified. Run from any directory.
"""

import argparse
from collections import Counter, defaultdict
import hashlib
import io
import json
import math
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
import shapefile
from shapely import STRtree, make_valid
from shapely.geometry import Point, mapping, shape


ROOT = Path(__file__).resolve().parents[1]
STATES = {1: "NSW", 2: "VIC", 3: "QLD", 4: "SA", 5: "WA", 6: "TAS", 7: "NT", 8: "ACT", 9: "Other territories"}
STATE_NAMES = dict(zip(
    ["New South Wales", "Victoria", "Queensland", "South Australia", "Western Australia", "Tasmania", "Northern Territory", "Australian Capital Territory"],
    ["NSW", "VIC", "QLD", "SA", "WA", "TAS", "NT", "ACT"],
))


def write(path, data, pretty=False):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, allow_nan=False,
                               indent=2 if pretty else None,
                               separators=None if pretty else (",", ":")) + "\n")


def numeric(value):
    if value is None or isinstance(value, str):
        raise ValueError(f"Expected a number, received {value!r}")
    value = float(value)
    assert math.isfinite(value)
    return value


def rounded_coords(value):
    if isinstance(value, (list, tuple)):
        return [rounded_coords(v) for v in value]
    return round(value, 5)


def build(source):
    used = set()
    audit = {"source_snapshot": "2026-09-10", "checks": {}, "issues": []}

    def read_json(relative):
        used.add(relative)
        return json.loads((source / relative).read_text())

    def workbook(relative):
        used.add(relative)
        return openpyxl.load_workbook(source / relative, read_only=True, data_only=True)

    # Region codes come from the original, unsimplified 2024 polygons.
    geo_path = "raw/tourism_boundaries/tra-tourism-regions-2024-shapefiles.zip"
    used.add(geo_path)
    with zipfile.ZipFile(source / geo_path) as archive:
        def member(suffix):
            return io.BytesIO(archive.read(next(n for n in archive.namelist() if n.endswith(suffix))))
        reader = shapefile.Reader(shp=member(".shp"), shx=member(".shx"), dbf=member(".dbf"))
        regions, geometries, geo_features, repairs = [], [], [], []
        for item in reader.iterShapeRecords():
            props = item.record.as_dict()
            region = {"region_code": props["TR_Code"], "region_name": props["TR_Name"], "state": STATES[int(props["State"])]}
            geom = shape(item.shape.__geo_interface__)
            if not geom.is_valid:
                repairs.append(region["region_code"])
                geom = make_valid(geom)
            assert geom.is_valid and not geom.is_empty
            regions.append(region)
            geometries.append(geom)
            # Display only: do not use simplified coastlines for point joins.
            simplified = geom.simplify(0.01, preserve_topology=True)
            out_geom = mapping(simplified)
            out_geom["coordinates"] = rounded_coords(out_geom["coordinates"])
            assert shape(out_geom).is_valid, region["region_code"]
            geo_features.append({"type": "Feature", "properties": region, "geometry": out_geom})
    assert len(regions) == len({r["region_code"] for r in regions}) == 76
    by_name = {r["region_name"]: r for r in regions}
    by_code = {r["region_code"]: r for r in regions}
    audit["checks"]["geometry_repairs"] = repairs
    write(ROOT / "data/tourism_regions_2024.geojson", {"type": "FeatureCollection", "features": geo_features})

    # Aggregate ABS SA2 population using the matching 2024 TRA allocation.
    w = workbook("raw/population/32180DS0003_2001-25.xlsx")
    population = {}
    for row in w["Table 1"].iter_rows(min_row=7, values_only=True):
        code = str(row[8])
        if len(code) == 9 and code.isdigit():
            assert code not in population
            population[code] = {2024: int(numeric(row[33])), 2025: int(numeric(row[34]))}
    w.close()
    w = workbook("raw/tourism_boundaries/tra-tourism-regions-2024-allocation.xlsx")
    allocated, region_pop = set(), defaultdict(lambda: {2024: 0, 2025: 0})
    for row in w.worksheets[0].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        code, tr = str(row[0]), str(row[3])
        assert code not in allocated and code in population and tr in by_code
        assert row[4] == by_code[tr]["region_name"]
        allocated.add(code)
        for year in (2024, 2025):
            region_pop[tr][year] += population[code][year]
    w.close()
    assert len(allocated) == 2452 and len(region_pop) == 76
    write(ROOT / "data/regional_population.json", [dict(r, population_2024=region_pop[r["region_code"]][2024], population_2025=region_pop[r["region_code"]][2025]) for r in regions])
    audit["checks"]["allocated_sa2_count"] = len(allocated)
    audit["checks"]["unallocated_population_rows"] = sorted(set(population) - allocated)

    # TRA publishes component sample sizes separately; preserve its thresholds.
    appendix_path = "raw/tourism_economy/appendix-1-change-in-direct-tourism-output.docx"
    used.add(appendix_path)
    with zipfile.ZipFile(source / appendix_path) as archive:
        document = ET.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    samples = []
    for row in document.findall(".//w:tbl/w:tr", ns):
        cells = ["".join(c.itertext()).strip() for c in row.findall("w:tc", ns)]
        if len(cells) != 8 or not all(c.isdigit() for c in cells[5:]):
            continue
        sizes = list(map(int, cells[5:]))
        samples.append({"region_name": cells[0], "state": cells[1].upper(), "domestic_day_sample": sizes[0],
                        "domestic_overnight_sample": sizes[1], "international_sample": sizes[2],
                        "below_source_threshold": sizes[0] < 1000 or sizes[1] < 500 or sizes[2] < 40})
    assert len(samples) == 76
    sample_by_name = {(s["state"], s["region_name"]): s for s in samples}
    write(ROOT / "data/regional_sample_sizes.json", samples)
    audit["checks"]["regions_below_any_sample_threshold"] = sum(s["below_source_threshold"] for s in samples)

    # No guessed name aliases: retain both unresolved economy records with null joins.
    economy, industry, unmatched = [], [], []
    fields = {6: "direct_gva_aud_million", 7: "indirect_gva_aud_million", 8: "total_gva_aud_million",
              10: "direct_grp_aud_million", 11: "indirect_grp_aud_million", 12: "total_grp_aud_million",
              14: "direct_filled_jobs_thousands", 15: "indirect_filled_jobs_thousands", 16: "total_filled_jobs_thousands",
              18: "tourism_consumption_aud_million"}
    for path in sorted((source / "raw/tourism_economy").glob("*-srtsa.xlsx")):
        relative = str(path.relative_to(source))
        w = workbook(relative)
        rows = list(w["Regional Summary"].values)
        name, state = rows[1][0], STATE_NAMES[rows[2][0]]
        region = by_name.get(name)
        assert region is None or region["state"] == state
        if region is None:
            unmatched.append({"economy_name": name, "state": state, "candidate_boundary_name": {"Capricorn": "Capricorn Coast", "Mackay Isaac": "Mackay"}.get(name), "status": "not joined; geographic equivalence unverified"})
        for col in (1, 2):
            year = rows[3][col]
            assert year in ("2023-24", "2024-25")
            record = {"region_name": name, "state": state, "financial_year": year,
                      "region_code": region["region_code"] if region else None,
                      "join_status": "exact_name_2024" if region else "unresolved_name",
                      "source_file": path.name}
            record.update({key: numeric(rows[row - 1][col]) for row, key in fields.items()})
            sample = sample_by_name.get((state, name))
            record["sample_caution"] = sample["below_source_threshold"] if sample else None
            for metric in ("gva_aud_million", "grp_aud_million", "filled_jobs_thousands"):
                assert math.isclose(record[f"direct_{metric}"] + record[f"indirect_{metric}"], record[f"total_{metric}"], rel_tol=1e-9)
            pop = region_pop[region["region_code"]][int("20" + year[-2:])] if region else None
            record["year_end_population"] = pop
            record["consumption_aud_per_resident"] = round(record["tourism_consumption_aud_million"] * 1e6 / pop, 2) if pop else None
            economy.append(record)
        jobs = list(w["Filled jobs"].values)
        header = next(i for i, row in enumerate(jobs) if row[1] == "Full -time")
        subtotal = 0
        for i in range(header + 1, len(jobs)):
            label = jobs[i][0]
            if label == "Direct tourism filled jobs":
                assert math.isclose(subtotal, numeric(jobs[i][3]), rel_tol=1e-8)
                assert math.isclose(subtotal, economy[-1]["direct_filled_jobs_thousands"], rel_tol=1e-8)
                break
            if not isinstance(jobs[i][3], (int, float)):
                continue
            full, part, total = map(numeric, jobs[i][1:4])
            assert math.isclose(full + part, total, rel_tol=1e-8)
            subtotal += total
            industry.append({"region_name": name, "state": state, "financial_year": "2024-25", "industry": label,
                             "full_time_thousands": full, "part_time_thousands": part, "total_thousands": total,
                             "source_file": path.name, "source_row": i + 1})
        else:
            raise ValueError(f"Missing jobs total in {path}")
        w.close()
    assert len(economy) == len({(r["state"], r["region_name"], r["financial_year"]) for r in economy}) == 152
    assert len(unmatched) == 2
    write(ROOT / "data/regional_economy.json", economy)
    write(ROOT / "data/regional_jobs_by_industry.json", industry)
    audit["issues"].append({"id": "economy_boundary_names", "records": unmatched,
                             "action": "Resolve from authoritative geographic metadata before mapping these two economy regions. Show missing data, never zero."})
    audit["checks"]["economy_geographic_matches"] = {"matched_regions": 74, "total_regions": 76}
    audit["issues"].append({"id": "sample_sizes", "action": "TRA flags component samples below 1000 domestic day, 500 domestic overnight, or 40 international. These are confidence cautions, not zero values or confidence intervals. Appendix Mornington Peninsula does not exact-match economy Peninsula, so that sample flag is null pending verification."})

    # OSM: object-level deduplication, access screen, then full-resolution spatial join.
    tree = STRtree(geometries)
    objects, category_counts, review = {}, {}, []
    for category, filename in [("lookout", "osm_lookouts_bbox.json"), ("garden", "osm_gardens_bbox.json"), ("visitor_centre", "osm_visitor_centres_bbox_retry.json")]:
        data = read_json("raw/attractions/" + filename)
        counts = Counter(input=len(data["elements"]))
        for element in data["elements"]:
            tags = element.get("tags", {})
            key = f'{element["type"]}/{element["id"]}'
            access = tags.get("access", "unknown")
            if access in ("private", "no"):
                counts["excluded_access"] += 1
                continue
            if category == "visitor_centre" and tags.get("tourism") != "information":
                counts["excluded_non_tourism_office"] += 1
                continue
            loc = element if element["type"] == "node" else element.get("center", {})
            if "lon" not in loc or "lat" not in loc:
                counts["missing_coordinate"] += 1
                review.append({"osm_id": key, "category": category, "reason": "missing_coordinate"})
                continue
            point = Point(loc["lon"], loc["lat"])
            hits = [int(i) for i in tree.query(point) if geometries[int(i)].covers(point)]
            if len(hits) != 1:
                reason = "outside_region_polygons" if not hits else "multiple_region_intersections"
                counts[reason] += 1
                review.append({"osm_id": key, "category": category, "reason": reason,
                               "longitude": loc["lon"], "latitude": loc["lat"], "region_candidates": [regions[i]["region_code"] for i in hits]})
                continue
            counts["accepted_memberships"] += 1
            if key not in objects:
                objects[key] = {"id": key, "lon": round(loc["lon"], 6), "lat": round(loc["lat"], 6),
                                "region_code": regions[hits[0]]["region_code"], "categories": [],
                                "coordinate_type": "node" if element["type"] == "node" else "bbox_center"}
                if tags.get("name"):
                    objects[key]["name"] = tags["name"]
                if access != "unknown":
                    objects[key]["access"] = access
            assert category not in objects[key]["categories"]
            objects[key]["categories"].append(category)
        assert counts["input"] == sum(v for k, v in counts.items() if k != "input")
        category_counts[category] = dict(counts, snapshot=data.get("osm3s", {}).get("timestamp_osm_base"))
    attractions = [objects[k] for k in sorted(objects)]
    # A review flag is not an instruction to merge distinct sites.
    names = defaultdict(list)
    for obj in attractions:
        if obj.get("name"):
            names[obj["name"].casefold().strip()].append(obj)
    possible_duplicates = []
    for group in names.values():
        for i, a in enumerate(group):
            for b in group[i + 1:]:
                distance_km = 111.2 * math.hypot((a["lon"] - b["lon"]) * math.cos(math.radians(a["lat"])), a["lat"] - b["lat"])
                if distance_km < 0.1:
                    possible_duplicates.append({"name": a["name"], "objects": [a["id"], b["id"]], "distance_m_approx": round(distance_km * 1000)})
    write(ROOT / "data/mapped_attractions.json", attractions)
    write(ROOT / "docs/attraction_review.json", {"unassigned": review, "same_name_within_100m": possible_duplicates}, pretty=True)
    audit["checks"]["osm"] = {"categories": category_counts, "unique_accepted_objects": len(attractions),
                                  "multi_category_objects": sum(len(o["categories"]) > 1 for o in attractions),
                                  "nearby_same_name_pairs_for_review": len(possible_duplicates)}
    audit["issues"].append({"id": "osm_coverage", "action": "Label counts as mapped OSM records, not unique attractions or a complete national inventory. Unknown access is not confirmed public access. Review same-name nearby pairs before site-level claims."})

    # Heritage is a separate registry, not a fourth mutually exclusive OSM category.
    heritage_raw = read_json("raw/attractions/national_heritage.geojson")["features"]
    heritage = {}
    for feature in heritage_raw:
        props = feature["properties"]
        if props["STATUS"] != "Listed place":
            continue
        key = str(props["PLACE_ID"])
        record = heritage.setdefault(key, {"place_id": key, "name": props["NAME"], "class": props["CLASS"], "states": [], "url": props["URL"]})
        assert record["class"] == props["CLASS"]
        if props["STATE"] not in record["states"]:
            record["states"].append(props["STATE"])
    for record in heritage.values():
        record["states"].sort()
        record["state_group"] = record["states"][0] if len(record["states"]) == 1 else "Multiple states"
    assert len(heritage) == 126
    write(ROOT / "data/listed_national_heritage.json", list(heritage.values()))
    audit["checks"]["heritage"] = {"input_features": len(heritage_raw), "listed_unique_places": len(heritage)}

    # Quarterly domestic tourism: null retains suppressed/unavailable values.
    domestic_path = "raw/domestic_visitors/tra-dots-march-quarter-2026-compared-to-march-quarter-2025.xlsx"
    w = workbook(domestic_path)
    sheets = {name: list(w[name].values) for name in ("Table 1a", "Table 1b", "Table 5b")}
    assert sheets["Table 1a"][3][2] == "March Quarter 2026"
    assert sheets["Table 1b"][3][5] == "March Quarter 2025"
    # Verify every affected metric, not just the trips block.
    for start in (6, 16, 26):
        for row in range(start, start + 9):
            assert sheets["Table 1b"][row - 1][5] == sheets["Table 1a"][row - 1][2]
    domestic = []
    for trip_type, sheet in [("overnight", "Table 1b"), ("day", "Table 5b")]:
        rows = sheets[sheet]
        for col, quarter in enumerate(["2025-Q1", "2025-Q2", "2025-Q3", "2025-Q4", "2026-Q1"], 1):
            for offset in range(9):
                state = rows[5 + offset][0]
                record = {"state": STATE_NAMES.get(state, "Australia" if state == "Australiab" else state), "quarter": quarter, "trip_type": trip_type, "source_sheet": sheet}
                for metric, row in [("trips_thousands", 6 + offset), ("expenditure_aud_million", (26 if trip_type == "overnight" else 16) + offset)]:
                    value = rows[row - 1][col]
                    record[metric] = numeric(value) if isinstance(value, (int, float)) else None
                    record[metric + "_cell"] = f"{chr(65 + col)}{row}"
                    if record[metric] is None:
                        record[metric + "_status"] = str(value)
                record["nights_thousands"] = numeric(rows[15 + offset][col]) if trip_type == "overnight" and isinstance(rows[15 + offset][col], (int, float)) else None
                domestic.append(record)
    w.close()
    assert len(domestic) == len({(d["state"], d["quarter"], d["trip_type"]) for d in domestic}) == 90
    write(ROOT / "data/domestic_state_quarters.json", domestic)
    audit["issues"].append({"id": "domestic_header_correction", "source": domestic_path, "cell": "Table 1b!F4", "original": "March Quarter 2025", "derived": "2026-Q1", "evidence": "All 27 numeric observations in the final column match Table 1a column C, explicitly March Quarter 2026. Table 5b also labels its final quarter March 2026. Original file unchanged."})

    cases = json.loads((ROOT / "scripts/case_study_extract.json").read_text())
    for filename in ("port_arthur_2024-25.pdf", "rbg_victoria_2024-25.pdf"):
        used.add("raw/case_studies/" + filename)
    write(ROOT / "data/community_case_studies.json", cases)
    audit["checks"]["case_study_metrics"] = len(cases)

    # File-level provenance: hashes identify exactly which original sources were read.
    url_by_path = {}
    for manifest in ("sources/download_manifest.json", "sources/boundaries_2024_check.json"):
        for entry in json.loads((source / manifest).read_text()):
            if entry.get("path"):
                url_by_path[entry["path"]] = entry.get("url")
    for entry in json.loads((source / "sources/case_study_downloads.json").read_text()):
        url_by_path["raw/case_studies/" + entry["file"]] = entry["url"]
    provenance = []
    for relative in sorted(used):
        data = (source / relative).read_bytes()
        url = url_by_path.get(relative)
        if relative.startswith("raw/attractions/osm_"):
            url = "https://overpass-api.de/api/interpreter"
        elif relative == "raw/attractions/national_heritage.geojson":
            url = "https://gis.environment.gov.au/gispubmap/rest/services/ogc_services/National_Heritage_List/MapServer/0"
        provenance.append({"path": relative, "url": url, "bytes": len(data), "sha256": hashlib.sha256(data).hexdigest()})
    write(ROOT / "docs/source_manifest.json", provenance, pretty=True)
    files = {p.name: {"bytes": p.stat().st_size, "sha256": hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted((ROOT / "data").glob("*.json"))}
    p = ROOT / "data/tourism_regions_2024.geojson"
    files[p.name] = {"bytes": p.stat().st_size, "sha256": hashlib.sha256(p.read_bytes()).hexdigest()}
    audit["outputs"] = files
    audit["checks"]["total_data_bytes"] = sum(f["bytes"] for f in files.values())
    write(ROOT / "docs/data_audit.json", audit, pretty=True)
    print(json.dumps(audit["checks"], indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, default=ROOT.parent / "Local Tourism Attractions Data")
    args = parser.parse_args()
    build(args.source_dir.resolve())
